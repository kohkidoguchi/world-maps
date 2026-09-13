"""Bank of Japan flow of funds - whom-to-whom detail for debt securities.

sjds.xlsx: "Debt securities holdings" by holder sector x issuer sector, quarterly stocks since 2015 Q1,
100 million yen (the only whom-to-whom table the BoJ publishes as a flat file; the full FF matrix is
already covered via the OECD financial accounts).
"""
from __future__ import annotations

import io
import logging
import re

import openpyxl

from ..common import Obs, http_get, fetched_at

log = logging.getLogger(__name__)
URL = "https://www.boj.or.jp/statistics/sj/sjds.xlsx"
SECTOR = {"GG": "GOV", "HN": "HH", "NFC": "NFC", "FC": "FIN", "NRES": "ROW", "CB": "CB", "DC": "FIN", "OFC": "FIN",
          "IPF": "FIN", "RES": "TOTAL", "ALL": "TOTAL"}


def fetch_all(store):
    try:
        raw = http_get(URL, ttl_hours=24 * 7)
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        header = next(r for r in rows if r and any(isinstance(c, str) and re.fullmatch(r"\d{4} Q[1-4]", c) for c in r))
        periods = {i: c.replace(" ", "-") for i, c in enumerate(header) if isinstance(c, str) and re.fullmatch(r"\d{4} Q[1-4]", c)}
        n = 0
        for r in rows:
            if not r or len(r) < 5:
                continue
            code = next((c for c in r[:6] if isinstance(c, str) and c.startswith("FB_")), None)
            m = re.fullmatch(r"FB_([A-Z]+)_DS_([A-Z]+)_XDC", code or "")
            if not m:
                continue
            holder, issuer = SECTOR.get(m.group(1)), SECTOR.get(m.group(2))
            if not holder or not issuer or issuer == "TOTAL":
                continue
            for i, p in periods.items():
                v = r[i] if i < len(r) else None
                try:
                    v = float(str(v).replace(",", ""))
                except (TypeError, ValueError):
                    continue
                store.add(Obs("BOJ_W2W", "JPN", holder, "FIN", entry="A", instrument="F3", cp_sector=issuer, measure="stock",
                              unit="XDC", unit_mult=8, freq="Q", period=p, value=v, note="debt securities holdings by issuer"))
                n += 1
        store.mark_source("BOJ_W2W", name="BoJ flow of funds - debt securities holdings by holder x issuer", url=URL, rows=n,
                          fetched_at=fetched_at(URL), license="BoJ terms (free re-use with attribution)")
        log.info("BoJ W2W: %d obs", n)
    except Exception as e:
        log.exception("BoJ failed")
        store.mark_source("BOJ_W2W", error=str(e)[:300])
